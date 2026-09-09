from __future__ import annotations

import csv
import hashlib
import io
import json
import mimetypes
import platform
import re
import sys
import zipfile
from collections import Counter
from collections.abc import Iterable
from pathlib import Path
from typing import Any, BinaryIO

from .database import utc_now
from .vtt import format_timestamp, to_srt, to_vtt

PACKAGE_VERSION = "1.1"
PRODUCT_VERSION = "1.3.0"
TURN_GAP_MS = 2_000

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def research_slug(value: str, fallback: str = "untitled") -> str:
    """Create concise, archive-friendly names without losing non-Latin scripts."""
    normalized = re.sub(r"[^\w]+", "_", value.strip().lower(), flags=re.UNICODE)
    normalized = re.sub(r"_+", "_", normalized).strip("_")
    return normalized[:80] or fallback


def timespan(start_ms: int, end_ms: int) -> str:
    return f"{format_timestamp(start_ms)}–{format_timestamp(end_ms)}"


def segment_id(recording: dict[str, Any], cue: dict[str, Any]) -> str:
    return f"{research_slug(recording['name'])}_{recording['id'][:8]}_seg_{int(cue['id']):06d}"


def review_summary(cues: list[dict[str, Any]]) -> dict[str, int | str]:
    verified = sum(cue["status"] == "verified" for cue in cues)
    unresolved = len(cues) - verified
    return {
        "total": len(cues),
        "verified": verified,
        "unresolved": unresolved,
        "label": "Human review complete" if cues and not unresolved else f"{verified} of {len(cues)} segments verified",
    }


def transcript_state(cues: list[dict[str, Any]]) -> str:
    return "verified" if cues and all(cue.get("status") == "verified" for cue in cues) else "draft"


def _clean_text(value: Any) -> str:
    return " ".join(str(value or "").split())


def build_turns(recording: dict[str, Any], cues: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Derive researcher-friendly speaker turns while retaining cue traceability."""
    turns: list[dict[str, Any]] = []
    for cue in cues:
        content = _clean_text(cue.get("final_text"))
        current = turns[-1] if turns else None
        same_speaker = current and current["speaker_id"] != "Unknown speaker" and current["speaker_id"] == cue.get("speaker", "Unknown speaker")
        small_gap = current and int(cue["start_ms"]) - int(current["end_ms"]) <= TURN_GAP_MS
        if current and same_speaker and small_gap:
            if content:
                current["content"] = " ".join(part for part in (current["content"], content) if part)
            current["end_ms"] = max(int(current["end_ms"]), int(cue["end_ms"]))
            current["segment_ids"].append(segment_id(recording, cue))
            current["cue_indexes"].append(int(cue["cue_index"]))
            current["review_statuses"].append(str(cue["status"]))
            continue

        turns.append(
            {
                "turn_id": "",  # assigned after grouping so IDs remain easy to read
                "start_ms": int(cue["start_ms"]),
                "end_ms": int(cue["end_ms"]),
                "speaker_id": cue.get("speaker", "Unknown speaker"),
                "speaker_role": speaker_role(str(cue.get("speaker", ""))),
                "content": content,
                "segment_ids": [segment_id(recording, cue)],
                "cue_indexes": [int(cue["cue_index"])],
                "review_statuses": [str(cue["status"])],
            }
        )
    prefix = f"{research_slug(recording['name'])}_{recording['id'][:8]}"
    for index, turn in enumerate(turns, start=1):
        turn["turn_id"] = f"{prefix}_turn_{index:04d}"
        turn["review_status"] = (
            "verified" if all(status == "verified" for status in turn.pop("review_statuses")) else "contains_unresolved"
        )
    return turns


def speaker_role(speaker: str) -> str:
    lowered = speaker.casefold()
    if any(token in lowered for token in ("moderator", "interviewer", "researcher", "facilitator")):
        return "Researcher / moderator"
    if speaker and lowered not in {"unknown", "unknown speaker"}:
        return "Participant"
    return "Unassigned"


def _initial_cue_states(audit_rows: list[dict[str, Any]]) -> dict[int, dict[str, Any]]:
    states: dict[int, dict[str, Any]] = {}
    for audit in audit_rows:
        cue_id = int(audit["cue_id"])
        if cue_id in states:
            continue
        before = audit.get("before")
        if before is None:
            raw = audit.get("before_json", "{}")
            try:
                before = json.loads(raw)
            except (json.JSONDecodeError, TypeError):
                before = {}
        states[cue_id] = before or {}
    return states


def _latest_human_actions(audit_rows: list[dict[str, Any]]) -> dict[int, dict[str, Any]]:
    actions: dict[int, dict[str, Any]] = {}
    for audit in audit_rows:
        try:
            cue_id = int(audit["cue_id"])
        except (KeyError, TypeError, ValueError):
            continue
        after = audit.get("after")
        if after is None:
            try:
                after = json.loads(audit.get("after_json", "{}"))
            except (json.JSONDecodeError, TypeError):
                after = {}
        if isinstance(after, dict) and after.get("status") in {"verified", "unclear"}:
            actions[cue_id] = audit
    return actions


def change_rows(
    recording: dict[str, Any], cues: list[dict[str, Any]], audit_rows: list[dict[str, Any]]
) -> list[dict[str, Any]]:
    initial = _initial_cue_states(audit_rows)
    latest_actions = _latest_human_actions(audit_rows)
    rows: list[dict[str, Any]] = []
    for cue in cues:
        before_state = initial.get(int(cue["id"]), {})
        before_text = _clean_text(cue.get("source_text"))
        after_text = _clean_text(cue.get("final_text"))
        before_speaker = str(before_state.get("speaker", cue.get("speaker", "Unknown speaker")))
        changes: list[str] = []
        if before_text != after_text:
            changes.append("wording")
        if before_speaker != str(cue.get("speaker", "Unknown speaker")):
            changes.append("speaker")
        if before_state and (
            int(before_state.get("start_ms", cue["start_ms"])) != int(cue["start_ms"])
            or int(before_state.get("end_ms", cue["end_ms"])) != int(cue["end_ms"])
        ):
            changes.append("timing")
        has_source = bool(before_text)
        action = latest_actions.get(int(cue["id"]))
        rows.append(
            {
                "recording_id": recording["id"],
                "recording_name": recording["name"],
                "segment_id": segment_id(recording, cue),
                "cue_index": cue["cue_index"],
                "start_time": format_timestamp(cue["start_ms"]),
                "end_time": format_timestamp(cue["end_ms"]),
                "speaker_before": before_speaker,
                "speaker_after": cue.get("speaker", "Unknown speaker"),
                "before_source_vtt": before_text,
                "after_current_text": after_text,
                "changed": ("Yes" if changes else "No") if has_source else "Not comparable — no source VTT",
                "change_type": ";".join(changes) or ("verification_only" if action else "not_human_reviewed"),
                "change_reason": ";".join(cue.get("ambiguity_reasons", [])) or "manual_review",
                "stt_a": cue.get("stt_a_text", ""),
                "stt_b": cue.get("stt_b_text", ""),
                "review_status": cue["status"],
                "reviewer": str(action.get("reviewer") or "Not recorded") if action else "",
                "reviewed_at": str(action.get("created_at") or "") if action else "",
            }
        )
    return rows


def segment_rows(recording: dict[str, Any], cues: list[dict[str, Any]]) -> list[dict[str, Any]]:
    turns = build_turns(recording, cues)
    turn_by_segment = {
        item: turn["turn_id"] for turn in turns for item in turn["segment_ids"]
    }
    return [
        {
            "recording_id": recording["id"],
            "recording_name": recording["name"],
            "segment_id": segment_id(recording, cue),
            "cue_index": cue["cue_index"],
            "turn_id": turn_by_segment.get(segment_id(recording, cue), ""),
            "start_time": format_timestamp(cue["start_ms"]),
            "end_time": format_timestamp(cue["end_ms"]),
            "start_ms": cue["start_ms"],
            "end_ms": cue["end_ms"],
            "speaker_id": cue.get("speaker", "Unknown speaker"),
            "speaker_role": speaker_role(str(cue.get("speaker", ""))),
            "content": _clean_text(cue.get("final_text")),
            "source_vtt_text": _clean_text(cue.get("source_text")),
            "stt_a": _clean_text(cue.get("stt_a_text")),
            "stt_b": _clean_text(cue.get("stt_b_text")),
            "confidence_a": cue.get("confidence_a"),
            "confidence_b": cue.get("confidence_b"),
            "ambiguity_score": cue.get("ambiguity_score", 0),
            "ambiguity_reasons": ";".join(cue.get("ambiguity_reasons", [])),
            "review_status": cue.get("status", "pending"),
            "reviewed_at": cue.get("updated_at", "") if cue.get("status") in {"verified", "unclear"} else "",
        }
        for cue in cues
    ]


def turn_rows(recording: dict[str, Any], cues: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return [
        {
            "recording_id": recording["id"],
            "recording_name": recording["name"],
            "turn_id": turn["turn_id"],
            "start_time": format_timestamp(turn["start_ms"]),
            "end_time": format_timestamp(turn["end_ms"]),
            "start_ms": turn["start_ms"],
            "end_ms": turn["end_ms"],
            "speaker_id": turn["speaker_id"],
            "speaker_role": turn["speaker_role"],
            "content": turn["content"],
            "source_segment_ids": ";".join(turn["segment_ids"]),
            "source_cue_indexes": ";".join(str(value) for value in turn["cue_indexes"]),
            "review_status": turn["review_status"],
        }
        for turn in build_turns(recording, cues)
    ]


def _csv_safe(value: Any) -> Any:
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def csv_bytes(rows: list[dict[str, Any]], fieldnames: list[str]) -> bytes:
    output = io.StringIO(newline="")
    writer = csv.DictWriter(output, fieldnames=fieldnames, extrasaction="ignore", lineterminator="\n")
    writer.writeheader()
    for row in rows:
        writer.writerow({field: _csv_safe(row.get(field, "")) for field in fieldnames})
    return output.getvalue().encode("utf-8-sig")


SEGMENT_FIELDS = [
    "recording_id", "recording_name", "segment_id", "cue_index", "turn_id", "start_time", "end_time",
    "start_ms", "end_ms", "speaker_id", "speaker_role", "content", "source_vtt_text", "stt_a", "stt_b",
    "confidence_a", "confidence_b", "ambiguity_score", "ambiguity_reasons", "review_status", "reviewed_at",
]

TURN_FIELDS = [
    "recording_id", "recording_name", "turn_id", "start_time", "end_time", "start_ms", "end_ms",
    "speaker_id", "speaker_role", "content", "source_segment_ids", "source_cue_indexes", "review_status",
]

CHANGE_FIELDS = [
    "recording_id", "recording_name", "segment_id", "cue_index", "start_time", "end_time",
    "speaker_before", "speaker_after", "before_source_vtt", "after_current_text", "changed", "change_type",
    "change_reason", "stt_a", "stt_b", "review_status", "reviewer", "reviewed_at",
]


def plain_transcript_bytes(recording: dict[str, Any], cues: list[dict[str, Any]]) -> bytes:
    summary = review_summary(cues)
    lines = [
        recording["name"],
        "=" * len(recording["name"]),
        f"Recording ID: {recording['id']}",
        f"Export schema version: {PACKAGE_VERSION}",
        f"Review status: {summary['label']}",
        "Transcription style: Verbatim wording; no automatic linguistic cleaning",
        "Identity review: Not assessed — review before sharing",
        "",
    ]
    for turn in build_turns(recording, cues):
        lines.append(f"{turn['speaker_id']}: {turn['content']} [{format_timestamp(turn['end_ms'])}]")
        lines.append("")
    return "\n".join(lines).encode("utf-8")


def analysis_docx_bytes(recording: dict[str, Any], cues: list[dict[str, Any]]) -> bytes:
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt

    document = Document()
    document.settings.odd_and_even_pages_header_footer = False
    section = document.sections[0]
    section.different_first_page_header_footer = False
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = document.styles
    _set_style_font(styles["Normal"], "Arial", 11)
    styles["Normal"].paragraph_format.space_before = Pt(0)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.25
    _set_style_font(styles["Heading 1"], "Arial", 16, bold=True, color="2E74B5")
    styles["Heading 1"].paragraph_format.space_before = Pt(18)
    styles["Heading 1"].paragraph_format.space_after = Pt(10)

    header = section.header.paragraphs[0]
    header.text = ""
    header_run = header.add_run("Transcript Review Studio  ·  Qualitative analysis transcript")
    _set_run_font(header_run, "Arial", 8.5, color="68717C", bold=True)
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.space_after = Pt(0)

    title = document.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(4)
    _set_run_font(title.add_run(recording["name"]), "Arial", 24, color="17365D", bold=True)
    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(16)
    _set_run_font(
        subtitle.add_run("Human-review workspace export · analysis-ready speaker turns"),
        "Arial",
        11,
        color="5F6B78",
        italic=True,
    )

    summary = review_summary(cues)
    metadata = [
        ("Recording ID", recording["id"]),
        ("Export schema version", PACKAGE_VERSION),
        ("Review status", str(summary["label"])),
        ("Segments", str(summary["total"])),
        ("Language", str(recording.get("metadata", {}).get("language", "Not recorded"))),
        ("Transcription style", "Verbatim wording; no automatic linguistic cleaning"),
        ("Identity review", "Not assessed — review before sharing"),
    ]
    table = document.add_table(rows=0, cols=2)
    table.style = "Light Shading Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    for label, value in metadata:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
        cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cells[0].paragraphs[0].runs[0].bold = True
        for cell in cells:
            cell.paragraphs[0].paragraph_format.space_before = Pt(0)
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
            cell.paragraphs[0].paragraph_format.line_spacing = 1.0
            for run in cell.paragraphs[0].runs:
                _set_run_font(run, "Arial", 9.5, color="24364B", bold=run.bold)
    _set_fixed_table_geometry(table, [2_376, 6_984], indent_dxa=120)

    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    document.add_heading("Transcript", level=1)
    for turn in build_turns(recording, cues):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.paragraph_format.line_spacing = 1.25
        paragraph.paragraph_format.keep_together = False
        paragraph.paragraph_format.keep_with_next = False
        paragraph.paragraph_format.widow_control = True
        speaker_run = paragraph.add_run(f"{turn['speaker_id']}: ")
        _set_run_font(speaker_run, "Arial", 11, color="125FB5", bold=True)
        _set_run_font(paragraph.add_run(turn["content"]), "Arial", 11, color="17191C")
        timestamp = paragraph.add_run(f" [{format_timestamp(turn['end_ms'])}]")
        _set_run_font(timestamp, "Arial", 8.5, color="6A737D", italic=True)

    footer = section.footer.paragraphs[0]
    footer.text = ""
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(0)
    footer.paragraph_format.space_after = Pt(0)
    _set_run_font(footer.add_run(f"Generated locally · Export schema {PACKAGE_VERSION} · Page "), "Arial", 8.5, color="68717C")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)

    document.core_properties.title = f"{recording['name']} — Qualitative Analysis Transcript"
    document.core_properties.subject = f"Research transcript — {transcript_state(cues)}"
    document.core_properties.author = "Transcript Review Studio"
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def nvivo_docx_bytes(recording: dict[str, Any], cues: list[dict[str, Any]]) -> bytes:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.shared import Inches, Pt

    document = Document()
    document.settings.odd_and_even_pages_header_footer = False
    section = document.sections[0]
    section.different_first_page_header_footer = False
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)
    _set_style_font(document.styles["Normal"], "Arial", 9.5)
    title = document.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    _set_run_font(title.add_run(f"{recording['name']} — NVivo transcript import"), "Arial", 20, color="17365D", bold=True)
    instruction = document.add_paragraph(
        "Import as transcript rows and map the columns to Timespan, Speaker, and Content. "
        "Stable Segment ID and Review Status may be mapped as custom fields."
    )
    instruction.paragraph_format.space_after = Pt(10)
    for run in instruction.runs:
        _set_run_font(run, "Arial", 9.5, color="5F6B78", italic=True)
    table = document.add_table(rows=1, cols=5)
    table.style = "Light Shading Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    headers = ["Timespan", "Speaker", "Content", "Segment ID", "Review Status"]
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
        cell.paragraphs[0].runs[0].bold = True
    for cue in cues:
        row = table.add_row().cells
        values = [
            timespan(cue["start_ms"], cue["end_ms"]),
            cue.get("speaker", "Unknown speaker"),
            _clean_text(cue.get("final_text")),
            segment_id(recording, cue),
            cue.get("status", "pending"),
        ]
        for cell, value in zip(row, values):
            cell.text = str(value)
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.paragraphs[0].paragraph_format.space_before = Pt(0)
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
            cell.paragraphs[0].paragraph_format.line_spacing = 1.1
            for run in cell.paragraphs[0].runs:
                _set_run_font(run, "Arial", 8.5, color="24364B", bold=run.bold)
    _set_fixed_table_geometry(table, [1_944, 1_656, 6_120, 3_096, 1_584], indent_dxa=120)
    table.rows[0]._tr.get_or_add_trPr().append(_repeat_table_header())
    for row in table.rows[1:]:
        row._tr.get_or_add_trPr().append(_keep_table_row_together())
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _repeat_table_header() -> Any:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    return repeat


def _keep_table_row_together() -> Any:
    """Prevent a single transcript segment from being split across Word pages."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    keep = OxmlElement("w:cantSplit")
    keep.set(qn("w:val"), "true")
    return keep


def _set_style_font(style: Any, name: str, size: float, *, bold: bool | None = None, color: str | None = None) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    style.font.name = name
    style.font.size = Pt(size)
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    if bold is not None:
        style.font.bold = bold
    if color:
        style.font.color.rgb = RGBColor.from_string(color)


def _set_run_font(
    run: Any,
    name: str,
    size: float,
    *,
    color: str | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def _set_fixed_table_geometry(table: Any, widths_dxa: list[int], *, indent_dxa: int) -> None:
    """Keep Word, LibreOffice, and Google importers on the same fixed column geometry."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    table_width = sum(widths_dxa)
    properties = table._tbl.tblPr
    width = properties.first_child_found_in("w:tblW")
    if width is None:
        width = OxmlElement("w:tblW")
        properties.append(width)
    width.set(qn("w:w"), str(table_width))
    width.set(qn("w:type"), "dxa")
    indent = properties.first_child_found_in("w:tblInd")
    if indent is None:
        indent = OxmlElement("w:tblInd")
        properties.append(indent)
    indent.set(qn("w:w"), str(indent_dxa))
    indent.set(qn("w:type"), "dxa")
    layout = properties.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        properties.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width_dxa in widths_dxa:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width_dxa))
        grid.append(column)
    for row in table.rows:
        for cell, width_dxa in zip(row.cells, widths_dxa):
            cell.width = width_dxa
            tc_width = cell._tc.get_or_add_tcPr().get_or_add_tcW()
            tc_width.set(qn("w:w"), str(width_dxa))
            tc_width.set(qn("w:type"), "dxa")
            margins = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcMar")
            if margins is None:
                margins = OxmlElement("w:tcMar")
                cell._tc.get_or_add_tcPr().append(margins)
            for edge, value in (("top", 80), ("bottom", 80), ("start", 120), ("end", 120)):
                margin = margins.find(qn(f"w:{edge}"))
                if margin is None:
                    margin = OxmlElement(f"w:{edge}")
                    margins.append(margin)
                margin.set(qn("w:w"), str(value))
                margin.set(qn("w:type"), "dxa")


def analysis_workbook_bytes(
    recording: dict[str, Any], cues: list[dict[str, Any]], audit_rows: list[dict[str, Any]]
) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    overview = workbook.active
    overview.title = "Overview"
    summary = review_summary(cues)
    overview_rows = [
        ("Qualitative Analysis Workbook", recording["name"]),
        ("Recording ID", recording["id"]),
        ("Export schema version", PACKAGE_VERSION),
        ("Review status", summary["label"]),
        ("Verified segments", summary["verified"]),
        ("Unresolved segments", summary["unresolved"]),
        ("Identity review", "Not assessed — review before sharing"),
        ("Structure", "Turns are derived from consecutive same-speaker segments separated by no more than 2 seconds."),
    ]
    for row in overview_rows:
        overview.append(row)
    for row in overview:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    overview.column_dimensions["A"].width = 24
    overview.column_dimensions["B"].width = 90
    overview["A1"].font = Font(bold=True, color="FFFFFF", size=14)
    overview["B1"].font = Font(bold=True, color="FFFFFF", size=14)
    overview["A1"].fill = overview["B1"].fill = PatternFill("solid", fgColor="087CFF")

    sheets = [
        ("Turns", turn_rows(recording, cues), TURN_FIELDS),
        ("Segments", segment_rows(recording, cues), SEGMENT_FIELDS),
        ("Change Log", change_rows(recording, cues, audit_rows), CHANGE_FIELDS),
    ]
    for title, rows, fields in sheets:
        sheet = workbook.create_sheet(title)
        sheet.append(fields)
        for row in rows:
            sheet.append([row.get(field, "") for field in fields])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="087CFF")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for index, field in enumerate(fields, start=1):
            content_width = 70 if field in {"content", "source_vtt_text", "before_source_vtt", "after_current_text", "stt_a", "stt_b"} else 22
            sheet.column_dimensions[get_column_letter(index)].width = content_width
        sheet.sheet_view.showGridLines = False

    workbook.properties.creator = "Transcript Review Studio"
    workbook.properties.title = f"{recording['name']} — Qualitative Analysis Workbook"
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def change_report_bytes(
    recording: dict[str, Any], cues: list[dict[str, Any]], audit_rows: list[dict[str, Any]]
) -> bytes:
    changes = change_rows(recording, cues, audit_rows)
    changed = [row for row in changes if row["changed"] == "Yes"]
    summary = review_summary(cues)
    lines = [
        f"# Transcript Change Report — {recording['name']}",
        "",
        f"- Recording ID: `{recording['id']}`",
        f"- Total segments: {summary['total']}",
        f"- Wording, speaker, or timing changes: {len(changed)}",
        f"- Human-verified segments: {summary['verified']}",
        f"- Unresolved or unclear segments: {summary['unresolved']}",
        "",
        "## Complete before/after list",
        "",
    ]
    for row in changes:
        lines.extend(
            [
                f"### {row['segment_id']} · {row['start_time']}–{row['end_time']}",
                "",
                f"- Speaker: {row['speaker_after']}",
                f"- Before (Source VTT): {row['before_source_vtt'] or '[no source text]'}",
                f"- After ({'Verified' if row['review_status'] == 'verified' else 'Current draft'}): {row['after_current_text'] or '[empty]'}",
                f"- Changed: {row['changed']} ({row['change_type']})",
                f"- Review status: {row['review_status']}",
                f"- Review evidence: STT A — {row['stt_a'] or '[no aligned text]'}; STT B — {row['stt_b'] or '[no aligned text]'}",
                f"- Flag reason: {row['change_reason']}",
                "",
            ]
        )
    return "\n".join(lines).encode("utf-8")


def review_actions_jsonl(audit_rows: list[dict[str, Any]]) -> bytes:
    lines: list[str] = []
    for row in audit_rows:
        item = {key: value for key, value in row.items() if key not in {"before_json", "after_json"}}
        for field in ("before", "after"):
            if field not in item:
                raw = row.get(f"{field}_json", "{}")
                try:
                    item[field] = json.loads(raw)
                except (json.JSONDecodeError, TypeError):
                    item[field] = {}
        lines.append(json.dumps(item, ensure_ascii=False, sort_keys=True))
    return (("\n".join(lines) + "\n") if lines else "").encode("utf-8")


def processing_provenance_bytes(recording: dict[str, Any], cues: list[dict[str, Any]]) -> bytes:
    metadata = recording.get("metadata", {})
    payload = {
        "product": "Transcript Review Studio",
        "product_version": PRODUCT_VERSION,
        "package_version": PACKAGE_VERSION,
        "recording_id": recording["id"],
        "recording_name": recording["name"],
        "processing_is_local": metadata.get("processing_is_local", True),
        "stt_model": metadata.get("model", "Not recorded"),
        "language_setting": metadata.get("language", "Not recorded"),
        "dual_pass": metadata.get("dual_pass", "Not recorded"),
        "named_speaker_tracks": metadata.get("named_speaker_tracks", []),
        "public_sample": metadata.get("public_sample"),
        "review": review_summary(cues),
        "runtime": {"python": sys.version.split()[0], "platform": platform.platform()},
        "statement": "Final text is exported exactly from the local human-verification workspace; no export-time wording cleanup is applied.",
    }
    return json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8")


def recordings_metadata_rows(
    payloads: list[dict[str, Any]], include_media: bool
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for payload in payloads:
        recording, cues = payload["recording"], payload["cues"]
        metadata = recording.get("metadata", {})
        summary = review_summary(cues)
        media_path = Path(recording["media_path"])
        rows.append(
            {
                "recording_id": recording["id"],
                "recording_name": recording["name"],
                "media_filename": media_path.name,
                "duration_ms": recording.get("duration_ms", 0),
                "media_type": metadata.get("media_type", "Not recorded"),
                "video_codec": metadata.get("video_codec", ""),
                "audio_codec": metadata.get("audio_codec", ""),
                "width": metadata.get("width", ""),
                "height": metadata.get("height", ""),
                "language": metadata.get("language", "Not recorded"),
                "stt_model": metadata.get("model", "Not recorded"),
                "dual_pass": metadata.get("dual_pass", "Not recorded"),
                "total_segments": summary["total"],
                "verified_segments": summary["verified"],
                "unresolved_segments": summary["unresolved"],
                "source_transcript_included": bool(recording.get("source_path")),
                "original_media_included": include_media,
                "identity_review_status": "Not assessed — review before sharing",
            }
        )
    return rows


RECORDING_FIELDS = [
    "recording_id", "recording_name", "media_filename", "duration_ms", "media_type", "video_codec",
    "audio_codec", "width", "height", "language", "stt_model", "dual_pass", "total_segments",
    "verified_segments", "unresolved_segments", "source_transcript_included", "original_media_included",
    "identity_review_status",
]

SPEAKER_FIELDS = [
    "recording_id", "recording_name", "speaker_id", "speaker_role", "segment_count", "speaking_duration_ms",
    "identity_status",
]


def speakers_metadata_rows(payloads: list[dict[str, Any]]) -> list[dict[str, Any]]:
    output: list[dict[str, Any]] = []
    for payload in payloads:
        recording, cues = payload["recording"], payload["cues"]
        counts = Counter(str(cue.get("speaker", "Unknown speaker")) for cue in cues)
        durations: Counter[str] = Counter()
        for cue in cues:
            durations[str(cue.get("speaker", "Unknown speaker"))] += int(cue["end_ms"]) - int(cue["start_ms"])
        for speaker in sorted(counts, key=str.casefold):
            output.append(
                {
                    "recording_id": recording["id"],
                    "recording_name": recording["name"],
                    "speaker_id": speaker,
                    "speaker_role": speaker_role(speaker),
                    "segment_count": counts[speaker],
                    "speaking_duration_ms": durations[speaker],
                    "identity_status": "Label as entered in Studio; pseudonymisation not automatically assessed",
                }
            )
    return output


def transcription_conventions_bytes() -> bytes:
    text = """# Transcription conventions

This package preserves the wording of `final_text` saved in Transcript Review Studio. Reading/structured formats normalize whitespace and combine eligible same-speaker segments. Raw cue JSON and the audit history retain saved text. No filler words, false starts, repetitions, grammatical features, or non-verbal notation are removed.

## Speaker turns and timing

- Each DOCX paragraph represents a derived speaker turn.
- Consecutive segments with the same assigned speaker are combined when the gap is 2 seconds or less. Unknown-speaker segments are not merged as if identity were known.
- The timestamp at the end of a DOCX/TXT paragraph marks the end of that turn and is compatible with common CAQDAS transcript workflows.
- Exact start and end times remain available in VTT, SRT, the NVivo DOCX table, CSV, XLSX, and JSON.
- Stable segment IDs connect every derived turn back to the reviewed cue and audit history.

## Notation

The application does not invent notation. If used by a researcher, bracketed forms such as `[laughs]`, `[overlapping speech]`, `[inaudible 00:12:34]`, or `[participant name]` are reproduced verbatim. The project team must document any additional conventions it adopts.

## Identity and confidentiality

Speaker labels and transcript content are exported as entered. The package does not claim that names, indirect identifiers, voices, or faces have been anonymised. Review identity risk, consent, and access conditions before sharing. Keep any identity crosswalk outside this package in appropriately restricted storage.
"""
    return text.encode("utf-8")


def data_dictionary_bytes() -> bytes:
    descriptions = {
        "recording_id": "Stable internal identifier for the recording.",
        "recording_name": "Human-readable recording title.",
        "segment_id": "Stable ID for a reviewed media-aligned cue.",
        "cue_index": "Display order of the cue within the recording.",
        "turn_id": "Derived speaker-turn identifier.",
        "start_time": "Cue or turn start in HH:MM:SS.mmm.",
        "end_time": "Cue or turn end in HH:MM:SS.mmm.",
        "start_ms": "Start offset from media origin in milliseconds.",
        "end_ms": "End offset from media origin in milliseconds.",
        "speaker_id": "Speaker label stored in the review workspace.",
        "speaker_role": "Role inferred only from explicit labels; otherwise Participant or Unassigned.",
        "content": "Final text saved in the human-verification workspace.",
        "source_vtt_text": "Text imported from the source VTT before review.",
        "before_source_vtt": "Source VTT wording before human review; blank when no source VTT was supplied.",
        "after_current_text": "Current workspace wording; consult review_status before treating it as verified.",
        "stt_a": "Aligned output from local STT pass A.",
        "stt_b": "Aligned output from independent local STT pass B.",
        "confidence_a": "Heuristic recognition score for STT pass A; not a calibrated probability.",
        "confidence_b": "Heuristic recognition score for STT pass B; not a calibrated probability.",
        "ambiguity_score": "Studio prioritisation score; higher values indicate more review evidence conflicts.",
        "ambiguity_reasons": "Semicolon-separated reasons the cue was queued for attention.",
        "review_status": "verified, pending, flagged, unclear, or contains_unresolved.",
        "reviewed_at": "Human review action time in ISO 8601 UTC; blank when not human reviewed.",
        "source_segment_ids": "Semicolon-separated segment IDs combined into a derived turn.",
        "source_cue_indexes": "Semicolon-separated cue display indexes combined into a turn.",
        "changed": "Whether wording, speaker, or timing differs from the starting state.",
        "change_type": "Semicolon-separated wording, speaker, timing, or verification_only classification.",
        "change_reason": "Original ambiguity reason(s), or manual_review.",
        "reviewer": "Reviewer identity; Not recorded when the local workspace has no named profile.",
    }
    rows = [
        {
            "dataset": dataset,
            "variable": field,
            "description": descriptions.get(field, "See README and transcription conventions."),
            "type": "string" if field not in {"cue_index", "start_ms", "end_ms", "ambiguity_score"} else "integer",
        }
        for dataset, fields in (("segments", SEGMENT_FIELDS), ("turns", TURN_FIELDS), ("change_log", CHANGE_FIELDS))
        for field in fields
    ]
    return csv_bytes(rows, ["dataset", "variable", "description", "type"])


def package_readme_bytes(project: dict[str, Any], payloads: list[dict[str, Any]], include_media: bool) -> bytes:
    total = sum(len(payload["cues"]) for payload in payloads)
    unresolved = sum(sum(cue["status"] != "verified" for cue in payload["cues"]) for payload in payloads)
    media_note = (
        "Original media is included. Treat this as a restricted package because voice and video may directly identify participants."
        if include_media
        else "Original media is excluded. Exact media filenames and technical metadata are retained so authorised researchers can relink local source files."
    )
    text = f"""# {project['name']} — Qualitative Analysis Package

- Package version: {PACKAGE_VERSION}
- Generated by: Transcript Review Studio {PRODUCT_VERSION}
- Recordings: {len(payloads)}
- Reviewed segments: {total - unresolved} of {total}
- Unresolved or unclear segments: {unresolved}

## Start here

- Use each `*_analysis_transcript.docx` for reading, coding, quotation, MAXQDA, or general CAQDAS import.
- Use each `*_nvivo_transcript.docx` to import media-linked rows into NVivo by mapping Timespan, Speaker, and Content.
- Use `*_analysis_workbook.xlsx` or the UTF-8 CSV files for framework analysis, mixed-method work, filtering, and reproducible scripts.
- Use VTT with the matching source media for timestamp-accurate playback in MAXQDA, ATLAS.ti, or other media-aware tools.
- Use files under `audit/` to inspect every before/after decision and the local processing provenance.

## Review status

A status of `verified` means a cue was explicitly saved as verified in the human-review workflow. Per-segment status is recorded in structured data, the NVivo table, and audit files. Reading transcripts carry a review summary. Caption filenames identify drafts; SRT has no review-status field, so keep the package metadata alongside it.

## Confidentiality

{media_note}

Transcript content and speaker labels are exported exactly as entered. Identity and disclosure risk have **not** been automatically assessed. Review direct and indirect identifiers, consent, copyright, and access conditions before sharing. Store any participant identity crosswalk separately.

## Integrity

`checksums.sha256` contains SHA-256 hashes for every payload and manifest file. `manifest.json` records package structure, review counts, file size, and media type. CSV files are UTF-8 with a byte-order mark for spreadsheet compatibility; potentially executable spreadsheet prefixes are escaped with a leading apostrophe in CSV only.

## Interoperability note

QDPX is not included because this workspace currently verifies source transcripts rather than creating analytic codes, cases, memos, or coded selections. QDPX becomes appropriate once those analytic objects exist and can be validated against the REFI-QDA schema.
"""
    return text.encode("utf-8")


def _sha256_bytes(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def _sha256_path(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        while chunk := stream.read(1024 * 1024):
            digest.update(chunk)
    return digest.hexdigest()


def _mime_for(path: str) -> str:
    suffix = Path(path).suffix.lower()
    explicit = {
        ".docx": DOCX_MIME,
        ".xlsx": XLSX_MIME,
        ".vtt": "text/vtt",
        ".srt": "application/x-subrip",
        ".jsonl": "application/x-ndjson",
        ".md": "text/markdown",
        ".csv": "text/csv",
    }
    return explicit.get(suffix) or mimetypes.guess_type(path)[0] or "application/octet-stream"


def _payload_entry(path: str, value: bytes | Path) -> dict[str, Any]:
    size = len(value) if isinstance(value, bytes) else value.stat().st_size
    digest = _sha256_bytes(value) if isinstance(value, bytes) else _sha256_path(value)
    return {"path": path, "size_bytes": size, "sha256": digest, "media_type": _mime_for(path)}


def write_project_package(
    output: BinaryIO,
    project: dict[str, Any],
    payloads: list[dict[str, Any]],
    *,
    include_media: bool = False,
) -> dict[str, Any]:
    files: dict[str, bytes | Path] = {}
    recording_manifest: list[dict[str, Any]] = []
    for payload in payloads:
        recording, cues, audit = payload["recording"], payload["cues"], payload["audit"]
        folder = f"{research_slug(recording['name'])}_{recording['id'][:8]}"
        base = research_slug(recording["name"])
        transcript_root = f"transcripts/{folder}"
        audit_root = f"audit/{folder}"
        files[f"{transcript_root}/{base}_analysis_transcript.docx"] = analysis_docx_bytes(recording, cues)
        files[f"{transcript_root}/{base}_nvivo_transcript.docx"] = nvivo_docx_bytes(recording, cues)
        state = transcript_state(cues)
        files[f"{transcript_root}/{base}_{state}.txt"] = plain_transcript_bytes(recording, cues)
        files[f"{transcript_root}/{base}_{state}.vtt"] = to_vtt(cues).encode("utf-8")
        files[f"{transcript_root}/{base}_{state}.srt"] = to_srt(cues).encode("utf-8")
        files[f"{transcript_root}/{base}_{state}.json"] = json.dumps(cues, ensure_ascii=False, indent=2).encode("utf-8")
        files[f"{transcript_root}/{base}_segments.csv"] = csv_bytes(segment_rows(recording, cues), SEGMENT_FIELDS)
        files[f"{transcript_root}/{base}_turns.csv"] = csv_bytes(turn_rows(recording, cues), TURN_FIELDS)
        files[f"{transcript_root}/{base}_analysis_workbook.xlsx"] = analysis_workbook_bytes(recording, cues, audit)
        files[f"{audit_root}/{base}_change_log.csv"] = csv_bytes(change_rows(recording, cues, audit), CHANGE_FIELDS)
        files[f"{audit_root}/{base}_change_report.md"] = change_report_bytes(recording, cues, audit)
        files[f"{audit_root}/{base}_review_actions.jsonl"] = review_actions_jsonl(audit)
        files[f"{audit_root}/{base}_processing_provenance.json"] = processing_provenance_bytes(recording, cues)
        source_path = Path(recording["source_path"]) if recording.get("source_path") else None
        if source_path and source_path.exists():
            files[f"source/{folder}/{research_slug(source_path.stem)}{source_path.suffix.lower()}"] = source_path
        media_path = Path(recording["media_path"])
        if include_media and media_path.exists():
            files[f"media/{folder}/{research_slug(media_path.stem)}{media_path.suffix.lower()}"] = media_path
        recording_manifest.append(
            {
                "recording_id": recording["id"],
                "recording_name": recording["name"],
                "folder": folder,
                "review": review_summary(cues),
                "speaker_count": len({cue.get("speaker", "Unknown speaker") for cue in cues}),
                "source_transcript_included": bool(source_path and source_path.exists()),
                "original_media_included": bool(include_media and media_path.exists()),
            }
        )

    files["README.md"] = package_readme_bytes(project, payloads, include_media)
    files["metadata/recordings.csv"] = csv_bytes(recordings_metadata_rows(payloads, include_media), RECORDING_FIELDS)
    files["metadata/speakers.csv"] = csv_bytes(speakers_metadata_rows(payloads), SPEAKER_FIELDS)
    files["metadata/data_dictionary.csv"] = data_dictionary_bytes()
    files["metadata/transcription_conventions.md"] = transcription_conventions_bytes()

    entries = [_payload_entry(path, value) for path, value in sorted(files.items())]
    manifest = {
        "package_type": "Qualitative Analysis Package",
        "package_version": PACKAGE_VERSION,
        "product": "Transcript Review Studio",
        "product_version": PRODUCT_VERSION,
        "exported_at": utc_now(),
        "project": {key: value for key, value in project.items() if key != "recordings"},
        "privacy": {
            "identity_review_status": "not_assessed",
            "original_media_included": include_media,
            "sharing_warning": "Review direct and indirect identifiers, consent, and access conditions before sharing.",
        },
        "recordings": recording_manifest,
        "files": entries,
    }
    manifest_bytes = json.dumps(manifest, ensure_ascii=False, indent=2).encode("utf-8")
    files["manifest.json"] = manifest_bytes
    checksum_lines = [
        f"{_payload_entry(path, value)['sha256']}  {path}" for path, value in sorted(files.items())
    ]
    files["checksums.sha256"] = ("\n".join(checksum_lines) + "\n").encode("utf-8")

    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=6) as archive:
        for path, value in sorted(files.items()):
            if isinstance(value, Path):
                compression = zipfile.ZIP_STORED if path.startswith("media/") else zipfile.ZIP_DEFLATED
                archive.write(value, path, compress_type=compression)
            else:
                archive.writestr(path, value)
    return manifest


def export_docx_bundle(
    format_name: str,
    recording: dict[str, Any],
    cues: list[dict[str, Any]],
    audit_rows: list[dict[str, Any]],
) -> tuple[bytes, str, str]:
    base = research_slug(recording["name"])
    if format_name == "docx":
        return analysis_docx_bytes(recording, cues), DOCX_MIME, f"{base}_analysis_transcript.docx"
    if format_name == "nvivo":
        return nvivo_docx_bytes(recording, cues), DOCX_MIME, f"{base}_nvivo_transcript.docx"
    if format_name == "xlsx":
        return analysis_workbook_bytes(recording, cues, audit_rows), XLSX_MIME, f"{base}_analysis_workbook.xlsx"
    if format_name == "txt":
        return plain_transcript_bytes(recording, cues), "text/plain; charset=utf-8", f"{base}_{transcript_state(cues)}.txt"
    raise ValueError(f"Unsupported qualitative export: {format_name}")


def all_speakers(cues: Iterable[dict[str, Any]]) -> list[str]:
    return sorted({str(cue.get("speaker", "Unknown speaker")) for cue in cues}, key=str.casefold)
