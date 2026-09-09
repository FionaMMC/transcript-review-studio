from __future__ import annotations

import hashlib
import io
import tempfile
import unittest
import zipfile
from pathlib import Path

from transcript_review.compare import create_aligned_cues, similarity, word_diff
from transcript_review.qualitative_export import (
    analysis_docx_bytes,
    analysis_workbook_bytes,
    build_turns,
    write_project_package,
)
from transcript_review.vtt import infer_speaker_prefix, parse_vtt, to_srt, to_vtt


class VttTests(unittest.TestCase):
    def test_cleared_editor_text_does_not_resurrect_source_in_captions(self) -> None:
        cue = {"start_ms": 0, "end_ms": 1000, "final_text": "", "source_text": "Removed hallucination", "speaker": "Unknown speaker"}
        self.assertNotIn("Removed hallucination", to_vtt([cue]))
        self.assertNotIn("Removed hallucination", to_srt([cue]))

    def test_malformed_timing_is_not_silently_discarded(self) -> None:
        with self.assertRaises(ValueError):
            parse_vtt("WEBVTT\n\n00:00:03.000 --> 00:00:01.000\nText")

    def test_round_trip_preserves_speaker_and_unicode(self) -> None:
        source = """WEBVTT

1
00:00:01.000 --> 00:00:03.500
<v Researcher>Hello &amp; welcome — 你好。
"""
        cues = parse_vtt(source)
        self.assertEqual(len(cues), 1)
        self.assertEqual(cues[0].speaker, "Researcher")
        self.assertEqual(cues[0].text, "Hello & welcome — 你好。")
        exported = to_vtt(
            [{"start_ms": cues[0].start_ms, "end_ms": cues[0].end_ms, "final_text": cues[0].text, "speaker": cues[0].speaker}]
        )
        self.assertIn("<v Researcher>Hello &amp; welcome — 你好。", exported)

    def test_srt_uses_comma_timestamps(self) -> None:
        output = to_srt([{"start_ms": 1000, "end_ms": 2500, "final_text": "Hello", "speaker": "Participant"}])
        self.assertIn("00:00:01,000 --> 00:00:02,500", output)
        self.assertIn("Participant: Hello", output)

    def test_plain_text_speaker_prefix_is_promoted(self) -> None:
        speaker, text = infer_speaker_prefix("P22: I can go first.")
        self.assertEqual(speaker, "P22")
        self.assertEqual(text, "I can go first.")
        cues = parse_vtt("WEBVTT\n\n00:00:00.000 --> 00:00:01.000\nResearcher: Thank you.")
        self.assertEqual(cues[0].speaker, "Researcher")
        self.assertEqual(cues[0].text, "Thank you.")

    def test_colon_in_ordinary_sentence_is_not_a_speaker(self) -> None:
        speaker, text = infer_speaker_prefix("Note: this is still transcript text.")
        self.assertIsNone(speaker)
        self.assertEqual(text, "Note: this is still transcript text.")


class ComparisonTests(unittest.TestCase):
    def test_timed_segment_does_not_leak_words_into_empty_window(self) -> None:
        from transcript_review.compare import text_for_window
        text, _, _ = text_for_window([{"start_ms": 0, "end_ms": 3000, "text": "Only later", "words": [
            {"start": 2.0, "end": 2.8, "word": " Only later"}
        ]}], 0, 1000)
        self.assertEqual(text, "")

    def test_disagreement_is_flagged(self) -> None:
        source = [{"start_ms": 1000, "end_ms": 3000, "text": "The prototype was calming.", "speaker": "P1"}]
        pass_a = [{"start_ms": 1000, "end_ms": 3000, "text": "The notification was coming.", "confidence": 0.51}]
        pass_b = [{"start_ms": 1000, "end_ms": 3000, "text": "The prototype was calling.", "confidence": 0.88}]
        cue = create_aligned_cues(source, pass_a, pass_b)[0]
        self.assertIn("word_disagreement", cue["ambiguity_reasons"])
        self.assertIn("low_confidence", cue["ambiguity_reasons"])
        self.assertEqual(cue["status"], "flagged")

    def test_diff_and_similarity(self) -> None:
        self.assertGreater(similarity("Hello, world", "hello world"), 0.9)
        kinds = {part["kind"] for part in word_diff("I said yes", "I set yes")}
        self.assertIn("changed", kinds)

    def test_word_timestamps_prevent_segment_duplication(self) -> None:
        source = [
            {"start_ms": 0, "end_ms": 1000, "text": "Hello there", "speaker": "P1"},
            {"start_ms": 1000, "end_ms": 2000, "text": "general Kenobi", "speaker": "P1"},
        ]
        long_segment = {
            "start_ms": 0,
            "end_ms": 2000,
            "text": "Hello there general Kenobi",
            "confidence": 0.9,
            "words": [
                {"start": 0.0, "end": 0.4, "word": " Hello", "probability": 0.96},
                {"start": 0.4, "end": 0.9, "word": " there", "probability": 0.95},
                {"start": 1.0, "end": 1.4, "word": " general", "probability": 0.94},
                {"start": 1.4, "end": 1.9, "word": " Kenobi", "probability": 0.93},
            ],
        }
        aligned = create_aligned_cues(source, [long_segment], [])
        self.assertEqual(aligned[0]["stt_a_text"], "Hello there")
        self.assertEqual(aligned[1]["stt_a_text"], "general Kenobi")

    def test_stt_words_outside_source_cues_create_missing_source_cue(self) -> None:
        source = [
            {"start_ms": 0, "end_ms": 1000, "text": "First answer.", "speaker": "P1"},
            {"start_ms": 3000, "end_ms": 4000, "text": "Last answer.", "speaker": "P1"},
        ]
        pass_a = [
            {
                "start_ms": 0,
                "end_ms": 4000,
                "text": "First answer. An omitted thought. Last answer.",
                "confidence": 0.9,
                "words": [
                    {"start": 0.1, "end": 0.4, "word": " First", "probability": 0.95},
                    {"start": 0.4, "end": 0.8, "word": " answer.", "probability": 0.95},
                    {"start": 1.4, "end": 1.7, "word": " An", "probability": 0.92},
                    {"start": 1.7, "end": 2.1, "word": " omitted", "probability": 0.9},
                    {"start": 2.1, "end": 2.5, "word": " thought.", "probability": 0.91},
                    {"start": 3.1, "end": 3.4, "word": " Last", "probability": 0.95},
                    {"start": 3.4, "end": 3.8, "word": " answer.", "probability": 0.95},
                ],
            }
        ]
        aligned = create_aligned_cues(source, pass_a, [])
        self.assertEqual(len(aligned), 3)
        omitted = aligned[1]
        self.assertEqual(omitted["source_text"], "")
        self.assertEqual(omitted["stt_a_text"], "An omitted thought.")
        self.assertIn("missing_source", omitted["ambiguity_reasons"])
        self.assertEqual(omitted["status"], "flagged")

    def test_second_pass_can_recover_speech_missing_from_source_and_first_pass(self) -> None:
        source = [{"start_ms": 0, "end_ms": 900, "text": "Opening.", "speaker": "P1"}]
        pass_a = [{"start_ms": 0, "end_ms": 900, "text": "Opening.", "confidence": 0.9}]
        pass_b = [{
            "start_ms": 0,
            "end_ms": 2_000,
            "text": "Opening. Missing answer.",
            "confidence": 0.9,
            "words": [
                {"start": 0.1, "end": 0.7, "word": " Opening.", "probability": 0.95},
                {"start": 1.2, "end": 1.5, "word": " Missing", "probability": 0.93},
                {"start": 1.5, "end": 1.9, "word": " answer.", "probability": 0.93},
            ],
        }]
        aligned = create_aligned_cues(source, pass_a, pass_b)
        self.assertEqual(len(aligned), 2)
        self.assertEqual(aligned[1]["stt_b_text"], "Missing answer.")
        self.assertIn("missing_source", aligned[1]["ambiguity_reasons"])


class QualitativeExportTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        root = Path(self.temporary.name)
        self.media = root / "interview.mp4"
        self.source = root / "interview.vtt"
        self.media.write_bytes(b"not-a-real-video-but-valid-for-package-testing")
        self.source.write_text("WEBVTT\n", encoding="utf-8")
        self.recording = {
            "id": "recording12345678",
            "name": "CH Focus Group",
            "media_path": str(self.media),
            "source_path": str(self.source),
            "duration_ms": 9_000,
            "metadata": {
                "media_type": "video",
                "video_codec": "h264",
                "audio_codec": "aac",
                "width": 1280,
                "height": 720,
                "language": "en",
                "model": "test-model",
                "dual_pass": True,
                "processing_is_local": True,
            },
        }
        self.cues = [
            {
                "id": 41,
                "cue_index": 1,
                "start_ms": 0,
                "end_ms": 2_000,
                "source_text": "Hello world",
                "stt_a_text": "Hello world",
                "stt_b_text": "Hello word",
                "final_text": "Hello world",
                "speaker": "Moderator",
                "confidence_a": 0.95,
                "confidence_b": 0.8,
                "ambiguity_score": 20,
                "ambiguity_reasons": ["word_disagreement"],
                "status": "verified",
                "updated_at": "2026-08-28T01:00:00+00:00",
            },
            {
                "id": 42,
                "cue_index": 2,
                "start_ms": 2_300,
                "end_ms": 4_000,
                "source_text": "How are you?",
                "stt_a_text": "How are you?",
                "stt_b_text": "How are you?",
                "final_text": "How are you?",
                "speaker": "Moderator",
                "confidence_a": 0.99,
                "confidence_b": 0.98,
                "ambiguity_score": 0,
                "ambiguity_reasons": [],
                "status": "verified",
                "updated_at": "2026-08-28T01:00:04+00:00",
            },
            {
                "id": 43,
                "cue_index": 3,
                "start_ms": 4_200,
                "end_ms": 6_000,
                "source_text": "I am find.",
                "stt_a_text": "I am fine.",
                "stt_b_text": "I am fine.",
                "final_text": "I am fine.",
                "speaker": "P01",
                "confidence_a": 0.97,
                "confidence_b": 0.97,
                "ambiguity_score": 60,
                "ambiguity_reasons": ["word_disagreement"],
                "status": "verified",
                "updated_at": "2026-08-28T01:00:06+00:00",
            },
        ]
        self.audit = [
            {
                "id": 1,
                "cue_id": 43,
                "action": "update",
                "before_json": '{"speaker":"P01","start_ms":4200,"end_ms":6000}',
                "after_json": '{"speaker":"P01","start_ms":4200,"end_ms":6000}',
                "created_at": "2026-08-28T01:00:06+00:00",
            }
        ]

    def tearDown(self) -> None:
        self.temporary.cleanup()

    def test_adjacent_same_speaker_cues_form_one_turn(self) -> None:
        turns = build_turns(self.recording, self.cues)
        self.assertEqual(len(turns), 2)
        self.assertEqual(turns[0]["content"], "Hello world How are you?")
        self.assertEqual(len(turns[0]["segment_ids"]), 2)

    def test_docx_and_xlsx_are_valid_office_files(self) -> None:
        from docx import Document
        from openpyxl import load_workbook

        document = Document(io.BytesIO(analysis_docx_bytes(self.recording, self.cues)))
        self.assertIn("CH Focus Group", "\n".join(paragraph.text for paragraph in document.paragraphs))
        workbook = load_workbook(io.BytesIO(analysis_workbook_bytes(self.recording, self.cues, self.audit)))
        self.assertEqual(workbook.sheetnames, ["Overview", "Turns", "Segments", "Change Log"])
        self.assertEqual(workbook["Segments"].max_row, 4)

    def test_package_has_documentation_integrity_and_optional_media(self) -> None:
        project = {"id": "project1", "name": "Research Interviews", "recordings": []}
        payloads = [{"recording": self.recording, "cues": self.cues, "audit": self.audit}]
        output = io.BytesIO()
        write_project_package(output, project, payloads)
        with zipfile.ZipFile(io.BytesIO(output.getvalue())) as archive:
            names = set(archive.namelist())
            self.assertIn("README.md", names)
            self.assertIn("manifest.json", names)
            self.assertIn("checksums.sha256", names)
            self.assertIn("metadata/data_dictionary.csv", names)
            self.assertFalse(any(name.startswith("media/") for name in names))
            checksums = archive.read("checksums.sha256").decode("utf-8").splitlines()
            for line in checksums:
                expected, path = line.split("  ", 1)
                self.assertEqual(hashlib.sha256(archive.read(path)).hexdigest(), expected)

        restricted = io.BytesIO()
        write_project_package(restricted, project, payloads, include_media=True)
        with zipfile.ZipFile(io.BytesIO(restricted.getvalue())) as archive:
            self.assertTrue(any(name.startswith("media/") for name in archive.namelist()))


if __name__ == "__main__":
    unittest.main()
